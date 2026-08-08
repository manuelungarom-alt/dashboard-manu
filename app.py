import os
import re
from io import BytesIO
from datetime import datetime, timedelta

import openpyxl
import dropbox
from flask import Flask, jsonify, request, send_from_directory

app = Flask(__name__, static_folder="static", static_url_path="")

DROPBOX_REFRESH_TOKEN = os.environ["DROPBOX_REFRESH_TOKEN"]
DROPBOX_APP_KEY = os.environ["DROPBOX_APP_KEY"]
DROPBOX_APP_SECRET = os.environ["DROPBOX_APP_SECRET"]
EXCEL_PATH = os.environ.get("EXCEL_PATH", "/Banco_Oficial_Manu_v7.xlsx")

dbx = dropbox.Dropbox(
    oauth2_refresh_token=DROPBOX_REFRESH_TOKEN,
    app_key=DROPBOX_APP_KEY,
    app_secret=DROPBOX_APP_SECRET,
)

# Mismos rangos que usa el bot de Telegram (deben mantenerse sincronizados a mano
# si el Excel se reestructura). Ver bot.py para el detalle de cómo se calcularon.
GASTO_RANGES = {
    "2026-08": (393, 639), "2026-09": (647, 796), "2026-10": (804, 953),
    "2026-11": (961, 1110), "2026-12": (1118, 1267),
}
INGRESO_RANGES = {
    "2026-08": (355, 504), "2026-09": (509, 953), "2026-10": (958, 1281),
    "2026-11": (1286, 1435), "2026-12": (1440, 1589),
}

CATEGORIAS = [
    "Vivienda", "Supermercado", "Auto", "Transporte", "Uber", "UADE", "UAI",
    "Brensolma", "Salidas", "Suscripciones", "Tecnologia", "Salud", "Seguros",
    "Iglesia", "Ofrenda", "Elim", "Familia", "Mama y Papa", "Tarjetas",
    "Ahorro", "Futbol", "Boludcompras", "Vacaciones", "Impuestos",
    "Monotributo", "Otros", "Kioscos", "Campamentos", "Dates", "Alquiler",
    "Casa/Deco",
]

MESES_ES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
            "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def mes_sheet_name(dt: datetime) -> str:
    return dt.strftime("%Y-%m")


def descargar_excel():
    _, res = dbx.files_download(EXCEL_PATH)
    return openpyxl.load_workbook(BytesIO(res.content), data_only=False)


def extraer_persona(motivo: str) -> tuple[str, str]:
    """Separa el tag [Nombre] del motivo, si esta presente."""
    m = re.search(r"\[(\w+)\]\s*$", motivo or "")
    if m:
        limpio = motivo[:m.start()].strip()
        return limpio, m.group(1)
    return motivo or "", "—"


def construir_resumen_mes(sheet: str):
    if sheet not in GASTO_RANGES or sheet not in INGRESO_RANGES:
        return None

    wb = descargar_excel()
    ws = wb[sheet.split("-")[0]]

    ini_g, fin_g = GASTO_RANGES[sheet]
    ini_i, fin_i = INGRESO_RANGES[sheet]

    transacciones = []
    por_categoria = {}
    egresos_total = 0.0

    for r in range(ini_g, fin_g + 1):
        fecha = ws.cell(row=r, column=1).value
        monto = ws.cell(row=r, column=2).value
        motivo = ws.cell(row=r, column=3).value
        categoria = ws.cell(row=r, column=4).value
        if not monto:
            continue
        desc, persona = extraer_persona(motivo)
        transacciones.append({
            "desc": desc or "(sin descripción)",
            "cat": categoria or "Sin categoría",
            "fecha": fecha.strftime("%d/%m/%Y") if fecha else "",
            "quien": persona,
            "monto": -abs(monto),
            "_fecha_raw": fecha,
        })
        if categoria:
            por_categoria[categoria] = por_categoria.get(categoria, 0) + abs(monto)
        egresos_total += abs(monto)

    ingresos_total = 0.0
    for r in range(ini_i, fin_i + 1):
        fecha = ws.cell(row=r, column=7).value
        monto = ws.cell(row=r, column=8).value
        motivo = ws.cell(row=r, column=9).value
        if not monto:
            continue
        desc, persona = extraer_persona(motivo)
        transacciones.append({
            "desc": desc or "(sin descripción)",
            "cat": "—",
            "fecha": fecha.strftime("%d/%m/%Y") if fecha else "",
            "quien": persona,
            "monto": abs(monto),
            "_fecha_raw": fecha,
        })
        ingresos_total += abs(monto)

    transacciones.sort(key=lambda t: t["_fecha_raw"] or datetime.min, reverse=True)
    for t in transacciones:
        del t["_fecha_raw"]

    cat_top = max(por_categoria.items(), key=lambda kv: kv[1]) if por_categoria else ("—", 0)

    anio, mes_num = map(int, sheet.split("-"))
    label = f"{MESES_ES[mes_num].upper()} {anio}"

    return {
        "periodo_label": label,
        "ingresos": round(ingresos_total, 2),
        "egresos": round(egresos_total, 2),
        "balance": round(ingresos_total - egresos_total, 2),
        "movimientos": len(transacciones),
        "categoria_top": {
            "nombre": cat_top[0],
            "monto": round(cat_top[1], 2),
            "pct": round(cat_top[1] / egresos_total * 100) if egresos_total else 0,
        },
        "por_categoria": [
            {"cat": c, "monto": round(m, 2)}
            for c, m in sorted(por_categoria.items(), key=lambda kv: -kv[1])
        ],
        "transacciones": transacciones,
    }


@app.route("/api/resumen")
def api_resumen():
    periodo = request.args.get("periodo", "mes")
    hoy = datetime.now()

    if periodo != "mes":
        return jsonify({"error": f"El período '{periodo}' todavía no está conectado — próximamente."}), 501

    sheet = request.args.get("mes") or mes_sheet_name(hoy)
    resumen = construir_resumen_mes(sheet)
    if resumen is None:
        return jsonify({"error": f"No hay datos preparados para {sheet}."}), 404
    return jsonify(resumen)


@app.route("/")
def index():
    return send_from_directory("static", "index.html")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)
